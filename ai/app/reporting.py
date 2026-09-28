"""Paginated A4 DOCX/PDF output from the same content model."""
import io
import json
import os
from pathlib import Path
from html import escape
from decimal import Decimal
from datetime import date

def money(value):
    return f'{Decimal(str(value)):,.0f}'.replace(',','.')+' VNĐ'

def document_bytes(title,blocks,format='pdf',sample=True):
    if format=='docx':
        from docx import Document
        from docx.shared import Mm,Pt,RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc=Document();s=doc.sections[0]
        s.page_width=Mm(210);s.page_height=Mm(297)
        s.left_margin=Mm(30);s.right_margin=Mm(20);s.top_margin=Mm(22);s.bottom_margin=Mm(20)
        style=doc.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(13)
        style.paragraph_format.space_after=Pt(6)
        for name in ['Title','Heading 1','Heading 2']:
            doc.styles[name].font.name='Times New Roman';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
        for item in doc.styles:
            for fonts in item.element.iter(qn('w:rFonts')):
                for attr in list(fonts.attrib):
                    if 'Theme' in attr or 'theme' in attr:del fonts.attrib[attr]
                for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
            for border in list(item.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        doc.styles['Title'].font.size=Pt(17)
        if blocks and blocks[0].get('nationalOnly'):
            p=doc.add_paragraph('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc')
            p.alignment=1
            for run in p.runs:run.bold=True
        elif blocks and blocks[0].get('letterhead'):
            table=doc.add_table(rows=1,cols=2)
            table.columns[0].width=Mm(66);table.columns[1].width=Mm(94)
            table.cell(0,0).text=blocks[0]['agency']+'\nSố: [chưa cấp]'
            table.cell(0,1).text='CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n[Địa danh, ngày ban hành chưa xác nhận]'
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.alignment=1
                        for run in p.runs:run.font.size=Pt(12)
            doc.add_paragraph('')
        title_paragraph=doc.add_paragraph(title,style='Title')
        if blocks and blocks[0].get('letterhead'):
            title_paragraph.alignment=1
            for run in title_paragraph.runs:run.font.size=Pt(14);run.bold=True
        for b in blocks:
            if b.get('heading'):doc.add_paragraph(b['heading'],style='Heading 2')
            for i,t in enumerate(b.get('text',[])):
                p=doc.add_paragraph(str(t))
                if b.get('keepTogether') and i<len(b['text'])-1:p.paragraph_format.keep_with_next=True
            if b.get('rows'):
                rows=b['rows'];table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Table Grid'
                for j,v in enumerate(rows[0]):table.rows[0].cells[j].text=str(v)
                repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for row in rows[1:]:
                    for c,v in zip(table.add_row().cells,row):c.text=str(v)
                for row in table.rows:
                    # Keep a complete evidence row on one page; repeat the header after a break.
                    no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
                    for c in row.cells:
                        for p in c.paragraphs:
                            for run in p.runs:run.font.size=Pt(10)
                for c in table.rows[0].cells:
                    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E8EDF3');c._tc.get_or_add_tcPr().append(shd)
        p=s.footer.paragraphs[0];p.alignment=2
        p.add_run('Trang ')
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
        out=io.BytesIO();doc.save(out);return out.getvalue()
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT, TA_CENTER
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font_dir=Path(os.getenv('APPRAISAL_FONT_DIR','C:/Windows/Fonts'))
    if 'AppraisalSerif' not in pdfmetrics.getRegisteredFontNames():
        regular=font_dir/'times.ttf';bold=font_dir/'timesbd.ttf'
        if not regular.exists():
            regular=Path('/usr/share/fonts/truetype/liberation2/LiberationSerif-Regular.ttf')
            bold=Path('/usr/share/fonts/truetype/liberation2/LiberationSerif-Bold.ttf')
        pdfmetrics.registerFont(TTFont('AppraisalSerif',str(regular)))
        pdfmetrics.registerFont(TTFont('AppraisalSerifBold',str(bold)))
    body=ParagraphStyle('Body',fontName='AppraisalSerif',fontSize=13,leading=17,spaceAfter=6)
    heading=ParagraphStyle('Heading',parent=body,fontName='AppraisalSerifBold',fontSize=13,spaceBefore=10,keepWithNext=True)
    title_style=ParagraphStyle('Title',parent=heading,fontSize=17,leading=22)
    small=ParagraphStyle('Small',parent=body,fontSize=9,leading=12)
    header_style=ParagraphStyle('Letterhead',parent=body,fontSize=12,leading=15,alignment=TA_CENTER)
    if blocks and blocks[0].get('letterhead'):
        title_style.fontSize=14;title_style.leading=18;title_style.alignment=TA_CENTER
    out=io.BytesIO()
    doc=SimpleDocTemplate(out,pagesize=(210*mm,297*mm),leftMargin=30*mm,rightMargin=20*mm,topMargin=22*mm,bottomMargin=20*mm)
    story=[Paragraph(escape(title),title_style),Spacer(1,8)]
    if blocks and blocks[0].get('nationalOnly'):
        story.insert(0,Paragraph('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM<br/>Độc lập - Tự do - Hạnh phúc',header_style))
    elif blocks and blocks[0].get('letterhead'):
        header=Table([[Paragraph(escape(blocks[0]['agency'])+'<br/>Số: [chưa cấp]',header_style),
                       Paragraph('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM<br/>Độc lập - Tự do - Hạnh phúc<br/>[Địa danh, ngày ban hành chưa xác nhận]',header_style)]],colWidths=[66*mm,94*mm])
        header.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),16)]))
        story.insert(0,header)
    for b in blocks:
        if b.get('keepTogether'):
            group=[]
            if b.get('heading'):group.append(Paragraph(escape(b['heading']),heading))
            group.extend(Paragraph(escape(str(t)).replace('\n','<br/>'),body) for t in b.get('text',[]))
            story.append(KeepTogether(group));continue
        if b.get('heading'):story.append(Paragraph(escape(b['heading']),heading))
        for t in b.get('text',[]):story.append(Paragraph(escape(str(t)).replace('\n','<br/>'),body))
        if b.get('rows'):
            rows=[[Paragraph(escape(str(v)),small) for v in row] for row in b['rows']]
            table=Table(rows,colWidths=[160*mm/len(rows[0])]*len(rows[0]),repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E8EDF3')),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
            story.extend([table,Spacer(1,8)])
    def footer(canvas,doc):
        canvas.setFont('AppraisalSerif',9)
        canvas.drawRightString(190*mm,12*mm,f'Trang {doc.page}')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()

def report_blocks(case,kind):
    from .draft_templates import draft_blocks
    from .legal import legal_overview
    draft=draft_blocks(case,kind)
    if draft:return draft
    run=case['runs'][-1] if case['runs'] else None
    docs={d['id']:d for d in case['documents']}
    titles={'report':'Báo cáo hỗ trợ kiểm tra BCNCKT','supplement':'Dự thảo yêu cầu bổ sung hồ sơ',
            'notice':'Dự thảo thông báo kết quả rà soát BCNCKT','decision':'Khung dự thảo quyết định phê duyệt dự án',
            'suspension':'Ghi chú rà soát trước khi tạm dừng thẩm định'}
    blocks=[{'heading':'Thông tin hồ sơ','text':[case['name'],'Địa phương công trình: '+case['province'],
        'Thời điểm đánh giá: '+date.fromisoformat(case['legalDate']).strftime('%d/%m/%Y'),'Cán bộ phụ trách: '+case['assignee'],
        'Phạm vi: hỗ trợ đọc, đối chiếu và kiểm tra; không thay thế kết luận thẩm định của cơ quan có thẩm quyền.']}]
    overview=legal_overview(case)
    blocks.append({'heading':'Chế độ pháp lý và phạm vi','text':[overview['profile']['label'],overview['profile']['reason'],
        overview['profile']['scopeLabel'],
        'Các khung biểu mẫu được triển khai cho CQ chuyên môn theo NĐ 217. Phạm vi khác hoặc hồ sơ lịch sử chỉ xuất báo cáo hỗ trợ, cần hoàn thiện biểu mẫu tương ứng.']})
    if not run or run['stale']:
        blocks.append({'heading':'Chưa có kết quả kiểm tra còn hiệu lực','text':['Tài liệu chưa được kiểm tra hoặc đã thay đổi. Cần chạy lại trước khi sử dụng kết quả.']})
        return titles[kind],blocks
    blocks.append({'heading':'Phương pháp và tình trạng','text':[run['aiStatus'],
        'Bộ quy tắc: '+run['ruleVersion'],'Các nhận xét và nguồn bên dưới cần chuyên viên kiểm tra. Không xác nhận chữ ký số hoặc tính đúng của số liệu hiện trường.']})
    if kind=='decision':
        blocks.append({'heading':'Điều kiện để hoàn thiện dự thảo','text':[
            'Chưa xác nhận đủ điều kiện phê duyệt. Chưa cấp số văn bản, ngày ban hành hoặc người ký.',
            'Điều 1: Nội dung dự án phải được đối chiếu với dữ liệu đã xác nhận.',
            'Điều 2: Trách nhiệm thực hiện và cơ quan liên quan do người có thẩm quyền xác định.',
            'Điều 3: Chỉ hoàn thiện hiệu lực và chữ ký sau quy trình phê duyệt hợp lệ.']})
    status={'consistent':'Số liệu/đầu vào khớp trong phạm vi kiểm tra','inconsistent':'Cần làm rõ khác biệt',
            'insufficient_evidence':'Chưa đủ dữ liệu','requires_specialist':'Cần chuyên viên xác nhận'}
    fs=run['findings']
    if kind=='supplement':fs=[f for f in fs if f['result']!='consistent']
    for i,f in enumerate(fs):
        lines=[status.get(f['result'],f['result'])+'. '+f['explanation']]
        if f['calculation']:lines.append('Phép tính: '+f['calculation'])
        if f['missingEvidence']:lines.append('Cần cung cấp/xác nhận: '+'; '.join(f['missingEvidence']))
        for reference in f.get('legalRefs',[]):lines.append('Căn cứ: '+reference['label']+' — '+reference['url'])
        for ref in f['sources'][:(4 if kind=='report' else 1)]:
            lines.append('Nguồn: '+docs.get(ref['documentId'],{}).get('name','Tài liệu')+' — '+ref['locator']+': '+ref['quote'])
        if len(f['sources'])>4:lines.append('Các nguồn còn lại có trong kết quả JSON và màn hình bằng chứng.')
        if f.get('review'):lines.append('Ý kiến chuyên viên: '+f['review']['note']+' — '+f['review']['actor'])
        blocks.append({'heading':f'{i+1}. '+f['title'],'text':lines})
    for i,note in enumerate(run.get('aiNotes',[])):
        blocks.append({'heading':f'Đề xuất AI cần xác minh {i+1}','text':[note['text'],
            *['Nguồn: '+docs.get(ref['documentId'],{}).get('name','Tài liệu')+' — '+ref['locator']+': '+ref['text'] for ref in note['sources']]]})
    if run.get('costComparison'):
        delta=run['costComparison']
        rows=[['Khoản mục','Lần đầu','Lần mới','Thay đổi']]+[[x['name'],money(x['before']),money(x['after']),money(x['difference'])] for x in delta['items']]
        blocks.append({'heading':'Đối chiếu chi phí','text':['Tiết kiệm ròng theo tổng: '+money(delta['netSavings'])+'. Chưa có kết luận cắt giảm chi phí.'],'rows':rows})
    blocks.append({'heading':'Kết luận của báo cáo hỗ trợ','text':[
        'Chưa thay thế thông báo kết quả thẩm định hoặc quyết định phê duyệt. Các nội dung chuyên ngành và dữ liệu chưa xác nhận cần được xử lý trước khi người có thẩm quyền kết luận.',
        'Rà soát nội bộ: '+(case['finalReview']['note'] if case.get('finalReview') else 'Chưa hoàn tất.')]})
    return titles[kind],blocks

def export_document(case,kind,format):
    if format=='json':
        return json.dumps(case,ensure_ascii=False,indent=2).encode('utf-8'),'application/json'
    title,blocks=report_blocks(case,kind)
    return document_bytes(title,blocks,format,case['sample']),('application/pdf' if format=='pdf' else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
